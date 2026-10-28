from datetime import timedelta
from unittest import skipUnless
from django.db import connection
from django.test import TestCase,Client,override_settings
from django.core.management import call_command
from django.core.exceptions import ValidationError
from django.utils import timezone
from rest_framework.test import APIClient
from .models import (
    AnalystGrant, Assignment, AuditEvent, Circle, Discussion, EmailOutbox,
    Enrollment, Invitation, Problem, Solve, Student, Tag, Topic,
    TopicCatalogVersion, User, Worksheet, WorksheetIssue, WorksheetItem,
)
from . import services as s

@override_settings(AUTH_PASSWORD_VALIDATORS=[],PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class DomainTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('seed_demo',verbosity=0)
        cls.admin=User.objects.get(email='admin@plusfolio.example')
        cls.teacher=User.objects.get(email='teacher@plusfolio.example')
        cls.mentor=User.objects.get(email='mentor@plusfolio.example')
        cls.student=User.objects.get(email='student@plusfolio.example')
        cls.analyst=User.objects.get(email='analyst@plusfolio.example')
        cls.multi=User.objects.get(email='multi@plusfolio.example')
        cls.circle=Circle.objects.get(name='Математическая мастерская')
        cls.other=Circle.objects.get(name='Олимпиадный старт')
        cls.foreign=Circle.objects.get(name='Геометрия в деталях')
        cls.issue=cls.circle.issues.order_by('pk').first()
        cls.problem=cls.issue.worksheet.items.order_by('position').first().problem
    def client_for(self,user):
        client=APIClient();client.force_login(user);return client
    def test_student_sees_all_peer_submissions_but_no_solution(self):
        result=self.client_for(self.student).get(f'/api/circles/{self.circle.pk}/journal/{self.issue.pk}')
        self.assertEqual(result.status_code,200)
        self.assertGreater(len(result.data['people']),1)
        self.assertFalse(result.data['show_group_stats'])
        self.assertTrue(all('solution' not in p for p in result.data['problems']))
    def test_student_analytics_is_strictly_personal(self):
        result=self.client_for(self.student).get('/api/analytics')
        self.assertTrue(result.data['personal_only']);self.assertEqual(result.data['student_count'],1)
        self.assertLessEqual(result.data['solved_pairs'],result.data['assigned_pairs'])
    def test_student_cannot_read_foreign_circle(self):
        self.assertEqual(self.client_for(self.student).get(f'/api/circles/{self.foreign.pk}').status_code,404)
    def test_student_cannot_mark_or_create(self):
        client=self.client_for(self.student)
        self.assertEqual(client.post(f'/api/circles/{self.circle.pk}/problems/{self.problem.pk}/solve',{'value':True,'version':0,'student_id':self.student.student.pk},format='json').status_code,403)
        self.assertEqual(client.post('/api/circles',{'name':'Нельзя'},format='json').status_code,403)
    def test_analyst_can_read_solutions_in_scope(self):
        response=self.client_for(self.analyst).get('/api/materials/problem')
        self.assertGreater(len(response.data),0);self.assertTrue(all('solution' in p for p in response.data))
        self.assertEqual(self.client_for(self.analyst).get(f'/api/circles/{self.foreign.pk}').status_code,404)
    def test_analyst_cannot_write_result(self):
        with self.assertRaises(s.PermissionDenied):s.mark(self.analyst,self.circle,self.problem.pk,{'value':True,'version':0,'student_id':self.student.student.pk})
    def test_multi_roles_union(self):
        self.assertEqual(set(self.multi.roles),{'teacher','analyst'})
        self.assertTrue(s.can_analyse_group(self.multi,self.circle))
        self.assertFalse(s.can_write(self.multi,self.circle))
        self.assertTrue(s.can_write(self.multi,self.other))
    def test_no_group_analytics_via_student_teacher_role_without_link(self):
        self.student.roles.append('teacher');self.student.save();s.profiles(self.student)
        response=self.client_for(self.student).get('/api/analytics?circle='+str(self.circle.pk))
        self.assertEqual(response.data['student_count'],1)
        self.assertTrue(response.data['personal_only'])
    def test_shared_solve_independent_discussion(self):
        row=Solve.objects.filter(student=self.student.student,problem=self.problem).first()
        s.mark(self.teacher,self.circle,self.problem.pk,{'student_id':self.student.student.pk,'value':True,'version':row.version if row else 0})
        self.assertEqual(Solve.objects.filter(student=self.student.student,problem=self.problem).count(),1)
        self.assertTrue(Solve.objects.get(student=self.student.student,problem=self.problem).solved)
        self.assertFalse(Discussion.objects.filter(circle=self.other,problem=self.problem,discussed=True).exists())
    def test_second_teacher_can_revoke_first_mark(self):
        row=Solve.objects.filter(student=self.student.student,problem=self.problem).first()
        s.mark(self.mentor,self.circle,self.problem.pk,{'student_id':self.student.student.pk,'value':False,'version':row.version if row else 0})
        self.assertFalse(Solve.objects.get(student=self.student.student,problem=self.problem).solved)
    def test_optimistic_lock_and_real_boolean(self):
        row=Solve.objects.filter(student=self.student.student,problem=self.problem).first()
        version=row.version if row else 0
        s.mark(self.teacher,self.circle,self.problem.pk,{'student_id':self.student.student.pk,'value':True,'version':version})
        with self.assertRaises(s.Conflict):s.mark(self.teacher,self.circle,self.problem.pk,{'student_id':self.student.student.pk,'value':False,'version':version})
        with self.assertRaises(ValidationError):s.mark(self.teacher,self.circle,self.problem.pk,{'student_id':self.student.student.pk,'value':'false','version':version+1})
    def test_late_enrollment_gets_all_worksheets_once(self):
        new=User.objects.create_user('late@plusfolio.example','unused',name='Новый ученик',roles=['student']);s.profiles(new)
        s.enroll(self.teacher,self.circle,new.email);s.enroll(self.teacher,self.circle,new.email)
        self.assertEqual(new.student.assignments.filter(issue__circle=self.circle).count(),self.circle.issues.count())
        self.assertEqual(Enrollment.objects.filter(student=new.student,circle=self.circle,left_at=None).count(),1)
    def test_departure_preserves_history_and_other_membership(self):
        count=self.student.student.assignments.count();s.unenroll(self.teacher,self.circle,self.student.email)
        self.assertEqual(self.student.student.assignments.count(),count)
        self.assertFalse(s.circles_for(self.student).filter(pk=self.circle.pk).exists())
        self.assertTrue(s.circles_for(self.student).filter(pk=self.other.pk).exists())
    def test_rejoin_reuses_assignments(self):
        count=self.student.student.assignments.count();s.unenroll(self.teacher,self.circle,self.student.email);s.enroll(self.teacher,self.circle,self.student.email)
        self.assertEqual(self.student.student.assignments.count(),count)
        self.assertEqual(Enrollment.objects.filter(circle=self.circle,student=self.student.student).count(),2)
    def test_multiple_teachers_and_independent_revocation(self):
        self.assertEqual(self.circle.teaching.filter(ended_at=None).count(),2)
        s.assign_teacher(self.admin,self.circle,self.mentor.email,True)
        self.assertTrue(s.can_write(self.teacher,self.circle));self.assertFalse(s.can_write(self.mentor,self.circle))
    def test_duplicate_issue_rejected(self):
        with self.assertRaises(ValidationError):s.issue_worksheet(self.teacher,self.circle,self.issue.worksheet_id,timezone.localdate())
    def test_published_content_and_topics_protected(self):
        self.problem.title='Нельзя'
        with self.assertRaises(ValidationError):self.problem.save()
        with self.assertRaises(ValidationError):self.problem.topics.clear()
    def test_wrong_tag_type_rejected(self):
        tag=Tag.objects.filter(kind='circle').first();topic=Topic.objects.filter(parent__isnull=False,is_other=False,archived=False).first()
        with self.assertRaises(ValidationError):s.create_material(self.teacher,'problem',{'title':'Новая','statement':'a','solution':'b','topics':[str(topic.pk)],'tags':[tag.pk]})
    def test_topic_other_cycle_and_version(self):
        root=Topic.objects.get(parent=None);node=s.add_topic(self.admin,'Тестовый раздел',root.pk)
        self.assertEqual(node.children.filter(is_other=True).count(),1)
        other=node.children.get(is_other=True)
        with self.assertRaises(ValidationError):s.add_topic(self.admin,'Нельзя',other.pk)
        with self.assertRaises(ValidationError):s.update_topic(self.admin,other.pk,{'name':'Другое'})
        child=s.add_topic(self.admin,'Подраздел',node.pk)
        with self.assertRaises(ValidationError):s.update_topic(self.admin,node.pk,{'parent_id':str(child.pk)})
        self.assertGreater(TopicCatalogVersion.objects.count(),1)
    def test_student_cannot_edit_topics(self):
        self.assertEqual(self.client_for(self.student).post('/api/topics',{'name':'Нельзя'},format='json').status_code,403)
    def test_overlap_does_not_inflate_analytics(self):
        response=self.client_for(self.teacher).get('/api/analytics')
        self.assertLess(response.data['assigned_pairs'],Assignment.objects.filter(issue__circle__in=[self.circle,self.other]).count()*6)
        row=next(r for r in response.data['rows'] if r['problem_id']==str(self.problem.pk))
        self.assertEqual(row['recipients'],12)
    def test_invite_lower_role_and_admin_exception(self):
        with self.assertRaises(s.PermissionDenied):s.invite(self.teacher,'boss@plusfolio.example','admin')
        inv=s.invite(self.admin,'newadmin@plusfolio.example','admin');self.assertEqual(inv.role,'admin')
    def test_invite_accept_atomic_single_use_and_auto_assign(self):
        invitation=s.invite(self.teacher,'new@plusfolio.example','student',self.circle)
        import json
        token=json.loads(s.cipher().decrypt(invitation.letters.get().encrypted_payload.encode()))['token']
        user=s.accept_invitation(token,{'name':'Иванов Иван Иванович','graduation_year':2031,'password':'DemoPassword2026!'})
        self.assertEqual(user.student.assignments.count(),self.circle.issues.count())
        with self.assertRaises(ValidationError):s.accept_invitation(token,{'name':'Дубль','password':'DemoPassword2026!'})
    def test_revoked_inviter_access_blocks_pending_invitation(self):
        invitation=s.invite(self.teacher,'wait@plusfolio.example','student',self.circle)
        import json
        token=json.loads(s.cipher().decrypt(invitation.letters.get().encrypted_payload.encode()))['token']
        s.assign_teacher(self.admin,self.circle,self.teacher.email,True)
        with self.assertRaises(s.PermissionDenied):s.valid_invitation(token)
    def test_existing_user_accepts_added_role_only_after_login(self):
        invitation=s.invite(self.admin,self.student.email,'teacher',self.circle)
        import json
        token=json.loads(s.cipher().decrypt(invitation.letters.get().encrypted_payload.encode()))['token']
        with self.assertRaises(s.PermissionDenied):s.accept_invitation(token,{})
        user=s.accept_invitation(token,{},self.student)
        self.assertEqual(set(user.roles),{'student','teacher'})
    def test_student_mailbox_and_all_users_forbidden(self):
        client=self.client_for(self.student)
        self.assertEqual(client.get('/api/users').status_code,403)
        self.assertEqual(client.get('/api/mailbox').data,[])
    def test_login_requires_csrf(self):
        client=Client(enforce_csrf_checks=True);client.get('/api/session')
        credentials={'email':self.teacher.email,'password':'PlusfolioDemo2026!'}
        self.assertEqual(client.post('/api/login',credentials,content_type='application/json').status_code,403)
        self.assertEqual(client.post('/api/login',credentials,content_type='application/json',HTTP_X_CSRFTOKEN=client.cookies['csrftoken'].value).status_code,200)
    def test_last_admin_guard_ignores_inactive_admins_and_preserves_multi_role_admin(self):
        backup=User.objects.create_user('backup-admin@example.test','unused',roles=['analyst','admin'],is_active=False)
        client=self.client_for(self.admin)
        response=client.patch('/api/users',{'id':self.admin.pk,'roles':['teacher']},format='json')
        self.assertEqual(response.status_code,400)
        self.admin.refresh_from_db();self.assertIn('admin',self.admin.roles)
        backup.is_active=True;backup.save(update_fields=['is_active'])
        response=client.patch('/api/users',{'id':self.admin.pk,'roles':['teacher']},format='json')
        self.assertEqual(response.status_code,200)
        self.admin.refresh_from_db();self.assertEqual(self.admin.roles,['teacher'])
        response=self.client_for(backup).patch('/api/users',{'id':backup.pk,'roles':['analyst']},format='json')
        self.assertEqual(response.status_code,400)
        backup.refresh_from_db();self.assertIn('admin',backup.roles)

    def test_teacher_cannot_administer_roles(self):
        self.assertEqual(self.client_for(self.teacher).patch('/api/users',{'id':self.student.pk,'roles':['admin']},format='json').status_code,403)
    def test_solution_not_in_student_material_responses(self):
        client=self.client_for(self.student)
        for path in ['/api/materials/problem',f'/api/materials/worksheet/{self.issue.worksheet_id}']:
            result=client.get(path);self.assertNotIn('"solution"',result.content.decode())

    def test_analyst_invite_scope_does_not_expand_via_other_roles(self):
        # The second circle is taught by multi, but its tag is outside analyst's grant.
        self.other.tags.set(Tag.objects.filter(name='Региональный проект'))
        self.assertTrue(s.can_write(self.multi,self.other))
        self.assertFalse(s.invitation_allowed(self.multi,'teacher',self.other))
        self.assertTrue(s.invitation_allowed(self.multi,'student',self.other))
        self.assertTrue(s.invitation_allowed(self.multi,'teacher',self.circle))
    def test_analyst_cannot_remove_teacher(self):
        with self.assertRaises(s.PermissionDenied):s.assign_teacher(self.analyst,self.circle,self.teacher.email,True)
    def test_personal_analytics_survives_all_departures(self):
        before=self.client_for(self.student).get('/api/analytics').data['assigned_pairs']
        s.unenroll(self.teacher,self.circle,self.student.email)
        s.unenroll(self.teacher,self.other,self.student.email)
        result=self.client_for(self.student).get('/api/analytics').data
        self.assertEqual(result['assigned_pairs'],before)
        self.assertEqual(result['student_count'],1)
        self.assertTrue(result['personal_only'])

    def test_analyst_cannot_invite_into_not_launched_circle(self):
        circle=Circle.objects.create(name='Без преподавателя',created_by=self.admin)
        circle.tags.set(self.circle.tags.all())
        self.assertFalse(s.invitation_allowed(self.analyst,'student',circle))
        self.assertFalse(s.invitation_allowed(self.analyst,'teacher',circle))
        self.assertTrue(s.invitation_allowed(self.admin,'teacher',circle))

    def test_grouped_journal_selects_all_or_subset_and_deduplicates_states(self):
        client=self.client_for(self.teacher)
        result=client.get(f'/api/circles/{self.circle.pk}/journal?issues=all')
        self.assertEqual(result.status_code,200)
        self.assertEqual(len(result.data['blocks']),self.circle.issues.count())
        keys=[(row['student_id'],row['problem_id']) for row in result.data['solves']]
        self.assertEqual(len(keys),len(set(keys)))
        result=client.get(f'/api/circles/{self.circle.pk}/journal?issues={self.issue.pk},{self.issue.pk}')
        self.assertEqual([block['issue_id'] for block in result.data['blocks']],[self.issue.pk])

    def test_grouped_journal_rejects_foreign_and_invalid_selection(self):
        client=self.client_for(self.teacher)
        foreign_issue=self.foreign.issues.first()
        for suffix in ['', 'garbage', '0', '-1', f'{self.issue.pk},{foreign_issue.pk}']:
            response=client.get(f'/api/circles/{self.circle.pk}/journal',{'issues':suffix})
            self.assertIn(response.status_code,[400,404])
        self.assertEqual(client.get(f'/api/circles/{self.circle.pk}/journal/export',{'issues':str(foreign_issue.pk)}).status_code,404)
        self.assertEqual(client.get(f'/api/circles/{self.foreign.pk}/journal/export').status_code,404)

    def test_grouped_journal_keeps_global_results_in_each_worksheet_block(self):
        issue_ids=list(self.circle.issues.values_list('pk',flat=True))
        result=self.client_for(self.teacher).get(f'/api/circles/{self.circle.pk}/journal',{'issues':','.join(map(str,issue_ids))})
        all_ids=[problem['id'] for block in result.data['blocks'] for problem in block['problems']]
        repeated={pid for pid in all_ids if all_ids.count(pid)>1}
        self.assertTrue(repeated)
        pid=next(iter(repeated))
        state=Solve.objects.filter(student=self.student.student,problem_id=pid).first()
        s.mark(self.teacher,self.circle,pid,{'student_id':self.student.student.pk,'value':True,'version':state.version if state else 0})
        updated=self.client_for(self.teacher).get(f'/api/circles/{self.circle.pk}/journal?issues=all').data
        self.assertEqual(sum(row['student_id']==self.student.student.pk and row['problem_id']==pid and row['solved'] for row in updated['solves']),1)
        self.assertGreater(sum(any(task['id']==pid for task in block['problems']) for block in updated['blocks']),1)

    def test_xlsx_download_contains_worksheet_blocks_and_real_status_values(self):
        import io
        from openpyxl import load_workbook
        response=self.client_for(self.teacher).get(f'/api/circles/{self.circle.pk}/journal/export')
        self.assertEqual(response.status_code,200)
        self.assertIn('spreadsheetml',response['Content-Type'])
        self.assertIn('.xlsx',response['Content-Disposition'])
        self.assertEqual(response['Cache-Control'],'no-store')
        workbook=load_workbook(io.BytesIO(response.content))
        sheet=workbook['Журнал']
        self.assertEqual(sheet.freeze_panes,'B7')
        titles=' '.join(str(cell.value or '') for row in sheet.iter_rows(min_row=5,max_row=5) for cell in row)
        for issue in self.circle.issues.all():self.assertIn(issue.worksheet.title,titles)
        self.assertIn(self.student.name,[sheet.cell(row,1).value for row in range(8,sheet.max_row+1)])
        self.assertIn('Разобрана в кружке',sheet['A7'].value)
        self.assertEqual(sheet.max_row,7+self.circle.enrollments.filter(left_at=None).count())
        self.assertFalse(any(cell.data_type=='f' for row in sheet for cell in row))
        self.assertTrue(all(cell.value in ('Да','Нет') for row in sheet.iter_rows(min_row=8,min_col=2,max_col=5) for cell in row))

    def test_tsv_export_selection_search_and_formula_injection(self):
        import csv,io
        self.student.name='=HYPERLINK("https://example.test")';self.student.save()
        client=self.client_for(self.teacher)
        response=client.get(f'/api/circles/{self.circle.pk}/journal/export',{'file_format':'tsv','issues':str(self.issue.pk),'q':'HYPERLINK'})
        self.assertEqual(response.status_code,200)
        rows=list(csv.reader(io.StringIO(response.content.decode('utf-8-sig')),delimiter='\t'))
        self.assertEqual(len(rows),4)
        self.assertTrue(rows[3][0].startswith("'="))
        self.assertIn(self.issue.worksheet.title,rows[0][1])
        self.assertEqual(len(rows[1]),2+self.issue.worksheet.items.count())
        from openpyxl import load_workbook
        response=client.get(f'/api/circles/{self.circle.pk}/journal/export',{'q':'HYPERLINK','issues':str(self.issue.pk)})
        cell=load_workbook(io.BytesIO(response.content)).active['A8']
        self.assertEqual(cell.data_type,'s');self.assertEqual(cell.value,self.student.name)

    def test_student_grouped_download_has_no_solutions_emails_or_peer_analytics(self):
        import io
        from openpyxl import load_workbook
        client=self.client_for(self.student)
        journal=client.get(f'/api/circles/{self.circle.pk}/journal?issues=all').data
        self.assertGreater(len(journal['people']),1)
        self.assertFalse(journal['show_group_stats'])
        self.assertTrue(all('solution' not in problem for block in journal['blocks'] for problem in block['problems']))
        response=client.get(f'/api/circles/{self.circle.pk}/journal/export')
        self.assertEqual(response.status_code,200)
        sheet=load_workbook(io.BytesIO(response.content)).active
        flat=[cell.value for row in sheet for cell in row]
        self.assertNotIn('Решено',flat)
        self.assertNotIn(self.student.email,flat)
        self.assertNotIn(self.problem.solution,flat)
        expected=1+sum(issue.worksheet.items.count() for issue in self.circle.issues.all())
        self.assertEqual(sheet.max_column,expected)

    def test_empty_circle_export_and_unsupported_format(self):
        import io
        from openpyxl import load_workbook
        circle=Circle.objects.create(name='Пустой кружок',created_by=self.admin)
        client=self.client_for(self.admin)
        result=client.get(f'/api/circles/{circle.pk}/journal?issues=all')
        self.assertEqual(result.data['blocks'],[])
        response=client.get(f'/api/circles/{circle.pk}/journal/export')
        self.assertEqual(response.status_code,200)
        sheet=load_workbook(io.BytesIO(response.content)).active
        self.assertIn('Пустой кружок',sheet['A1'].value)
        response=client.get(f'/api/circles/{circle.pk}/journal/export',{'file_format':'pdf'})
        self.assertEqual(response.status_code,400)

    def test_analyst_export_is_read_only_and_in_scope(self):
        client=self.client_for(self.analyst)
        self.assertEqual(client.get(f'/api/circles/{self.circle.pk}/journal/export').status_code,200)
        self.assertEqual(client.get(f'/api/circles/{self.foreign.pk}/journal/export').status_code,404)
        self.assertEqual(client.post(f'/api/circles/{self.circle.pk}/journal/export',{},format='json').status_code,405)

    def test_multiple_selected_worksheets_exclude_unselected_third(self):
        extra=Worksheet.objects.exclude(pk__in=self.circle.issues.values('worksheet_id')).first()
        s.issue_worksheet(self.teacher,self.circle,extra.pk,timezone.localdate())
        chosen=list(self.circle.issues.exclude(worksheet=extra).values_list('pk',flat=True))
        response=self.client_for(self.teacher).get(f'/api/circles/{self.circle.pk}/journal',{'issues':','.join(map(str,chosen))})
        self.assertEqual({block['issue_id'] for block in response.data['blocks']},set(chosen))
        self.assertEqual(len(response.data['blocks']),2)

    def _reveal_solution(self,circle=None,user=None):
        circle=circle or self.circle;user=user or self.teacher
        row=Discussion.objects.filter(circle=circle,problem=self.problem).first()
        reviewed=s.mark(user,circle,self.problem.pk,{'value':True,'version':row.version if row else 0},discussion=True)
        return s.mark(user,circle,self.problem.pk,{'value':True,'version':reviewed['version']},solution_visibility=True)

    def test_student_and_analyst_solution_scopes_unite_and_revoke_independently(self):
        self.student.roles=['student','analyst'];self.student.save(update_fields=['roles'])
        tag=Tag.objects.create(kind='circle',name='Independent solution scope')
        self.foreign.tags.add(tag)
        grant=AnalystGrant.objects.create(user=self.student,tag=tag)
        analyst_ids=set(Problem.objects.filter(items__worksheet__issues__circle=self.foreign).values_list('pk',flat=True))
        problem_ids=list(Problem.objects.values_list('pk',flat=True))
        with self.assertNumQueries(1):
            self.assertEqual(s.visible_solution_ids(self.student,problem_ids),analyst_ids)
        self._reveal_solution()
        self.assertEqual(s.visible_solution_ids(self.student,problem_ids),analyst_ids|{self.problem.pk})
        self.assertEqual(s.visible_solution_ids(self.student,problem_ids,self.other),analyst_ids)
        grant.delete()
        self.assertEqual(s.visible_solution_ids(self.student,problem_ids),{self.problem.pk})
        s.unenroll(self.teacher,self.circle,self.student.email)
        self.assertEqual(s.visible_solution_ids(self.student,problem_ids),set())
        self.assertFalse(s.solution_allowed(self.student,self.problem))

    def test_solution_visibility_defaults_to_hidden_even_for_discussed_tasks(self):
        self.assertTrue(Discussion.objects.filter(discussed=True).exists())
        self.assertFalse(Discussion.objects.filter(solution_visible=True).exists())
        self.assertFalse(s.solution_allowed(self.student,self.problem,self.circle))

    def test_solution_cannot_be_revealed_before_discussion(self):
        row=Discussion.objects.filter(circle=self.circle,problem=self.problem).first()
        reviewed=s.mark(self.teacher,self.circle,self.problem.pk,{'value':False,'version':row.version if row else 0},discussion=True)
        response=self.client_for(self.teacher).post(f'/api/circles/{self.circle.pk}/problems/{self.problem.pk}/solution-visibility',{'value':True,'version':reviewed['version']},format='json')
        self.assertEqual(response.status_code,400)
        self.assertFalse(Discussion.objects.get(circle=self.circle,problem=self.problem).solution_visible)

    def test_revealed_solution_is_available_in_journal_bank_and_worksheet(self):
        self._reveal_solution()
        client=self.client_for(self.student)
        result=client.get(f'/api/circles/{self.circle.pk}/journal?issues=all').data
        copies=[problem for block in result['blocks'] for problem in block['problems'] if problem['id']==str(self.problem.pk)]
        self.assertGreater(len(copies),1)
        self.assertTrue(all(problem['solution']==self.problem.solution for problem in copies))
        self.assertTrue(next(state for state in result['discussions'] if state['problem_id']==str(self.problem.pk))['solution_visible'])
        detail=client.get(f'/api/materials/problem/{self.problem.pk}').data
        self.assertEqual(detail['solution'],self.problem.solution)
        sheet=client.get(f'/api/materials/worksheet/{self.issue.worksheet_id}').data
        self.assertEqual(next(p for p in sheet['problems'] if p['id']==str(self.problem.pk))['solution'],self.problem.solution)

    def test_solution_visibility_is_scoped_to_circle(self):
        self._reveal_solution()
        client=self.client_for(self.student)
        allowed=client.get(f'/api/materials/problem/{self.problem.pk}',{'circle':str(self.circle.pk)})
        hidden=client.get(f'/api/materials/problem/{self.problem.pk}',{'circle':str(self.other.pk)})
        self.assertEqual(allowed.data['solution'],self.problem.solution)
        self.assertNotIn('solution',hidden.data)
        journal=client.get(f'/api/circles/{self.other.pk}/journal?issues=all').data
        self.assertNotIn('solution',next(p for block in journal['blocks'] for p in block['problems'] if p['id']==str(self.problem.pk)))
        self.assertEqual(client.get(f'/api/materials/problem/{self.problem.pk}',{'circle':str(self.foreign.pk)}).status_code,404)

    def test_only_journal_editors_can_change_solution_visibility(self):
        row=Discussion.objects.get(circle=self.circle,problem=self.problem)
        path=f'/api/circles/{self.circle.pk}/problems/{self.problem.pk}/solution-visibility'
        for user in [self.student,self.analyst,self.multi]:
            response=self.client_for(user).post(path,{'value':True,'version':row.version},format='json')
            self.assertEqual(response.status_code,403)
        response=self.client_for(self.mentor).post(path,{'value':True,'version':row.version},format='json')
        self.assertEqual(response.status_code,200)
        response=self.client_for(self.admin).post(path,{'value':False,'version':response.data['version']},format='json')
        self.assertEqual(response.status_code,200)
        self.assertFalse(response.data['solution_visible'])

    def test_unmark_discussion_hides_solution_and_does_not_auto_reopen(self):
        result=self._reveal_solution()
        result=s.mark(self.teacher,self.circle,self.problem.pk,{'value':False,'version':result['version']},discussion=True)
        self.assertFalse(result['solution_visible'])
        result=s.mark(self.teacher,self.circle,self.problem.pk,{'value':True,'version':result['version']},discussion=True)
        self.assertFalse(result['solution_visible'])
        self.assertFalse(s.solution_allowed(self.student,self.problem,self.circle))
        self.assertTrue(AuditEvent.objects.filter(kind='solution.visibility',circle=self.circle).exists())

    def test_revoke_reveal_checks_shared_version_and_current_membership(self):
        result=self._reveal_solution()
        with self.assertRaises(s.Conflict):s.mark(self.mentor,self.circle,self.problem.pk,{'value':False,'version':result['version']-1},solution_visibility=True)
        self.assertTrue(s.solution_allowed(self.student,self.problem))
        s.unenroll(self.teacher,self.circle,self.student.email)
        self.assertFalse(s.solution_allowed(self.student,self.problem))
        self.assertNotIn('solution',self.client_for(self.student).get(f'/api/materials/problem/{self.problem.pk}').data)

    def test_database_rejects_visible_solution_without_discussion(self):
        from django.db import IntegrityError,transaction
        self._reveal_solution()
        with self.assertRaises(IntegrityError):
            with transaction.atomic():Discussion.objects.filter(circle=self.circle,problem=self.problem).update(discussed=False)

    def test_revealed_solutions_do_not_enter_journal_exports(self):
        import io,csv
        from openpyxl import load_workbook
        self._reveal_solution()
        client=self.client_for(self.student)
        response=client.get(f'/api/circles/{self.circle.pk}/journal/export')
        sheet=load_workbook(io.BytesIO(response.content)).active
        self.assertNotIn(self.problem.solution,[cell.value for row in sheet for cell in row])
        response=client.get(f'/api/circles/{self.circle.pk}/journal/export',{'file_format':'tsv'})
        self.assertNotIn(self.problem.solution,response.content.decode('utf-8-sig'))

    def test_revoke_reveal_removes_solution_from_next_detail_request(self):
        opened=self._reveal_solution()
        client=self.client_for(self.student)
        path=f'/api/materials/problem/{self.problem.pk}?circle={self.circle.pk}'
        self.assertIn('solution',client.get(path).data)
        s.mark(self.teacher,self.circle,self.problem.pk,{'value':False,'version':opened['version']},solution_visibility=True)
        self.assertNotIn('solution',client.get(path).data)
        self.assertTrue(Discussion.objects.get(circle=self.circle,problem=self.problem).discussed)

    def test_student_year_is_optional_and_does_not_disclose_email(self):
        response=self.client_for(self.student).get(f'/api/circles/{self.circle.pk}/members')
        person=next(p for p in response.data if p['id']==self.student.student.pk)
        self.assertIsNone(person['graduation_year']);self.assertIsNone(person['current_grade'])
        self.assertEqual(person['grade_label'],'Класс не указан')
        self.assertNotIn('email',person)

    def test_student_year_is_shared_between_circles_and_journal_with_audit(self):
        from datetime import date
        from unittest.mock import patch
        client=self.client_for(self.teacher)
        solves=Solve.objects.filter(student=self.student.student).count()
        assignments=self.student.student.assignments.count()
        with patch('core.school.timezone.localdate',return_value=date(2026,10,21)):
            result=client.patch(f'/api/circles/{self.circle.pk}/members/{self.student.student.pk}',{'graduation_year':2031},format='json')
            self.assertEqual(result.status_code,200);self.assertEqual(result.data['current_grade'],7)
            for circle in [self.circle,self.other]:
                journal=self.client_for(self.student).get(f'/api/circles/{circle.pk}/journal').data
                person=next(p for p in journal['people'] if p['id']==self.student.student.pk)
                self.assertEqual(person['graduation_year'],2031);self.assertEqual(person['grade_label'],'7 класс')
            self.assertEqual(self.client_for(self.student).get('/api/session').json()['user']['current_grade'],7)
        event=AuditEvent.objects.get(kind='student.update',circle=self.circle)
        self.assertEqual(event.actor,self.teacher)
        self.assertEqual(event.payload,{'student_id':self.student.student.pk,'old_graduation_year':None,'graduation_year':2031})
        self.assertEqual(Solve.objects.filter(student=self.student.student).count(),solves)
        self.assertEqual(self.student.student.assignments.count(),assignments)

    def test_both_teachers_and_admin_can_correct_and_clear_student_year(self):
        path=f'/api/circles/{self.circle.pk}/members/{self.student.student.pk}'
        for user,year in [(self.teacher,2031),(self.mentor,2032),(self.admin,None)]:
            result=self.client_for(user).patch(path,{'graduation_year':year},format='json')
            self.assertEqual(result.status_code,200);self.assertEqual(result.data['graduation_year'],year)
        self.assertEqual(result.data['grade_label'],'Класс не указан')
        self.assertEqual(AuditEvent.objects.filter(kind='student.update').count(),3)
        self.client_for(self.admin).patch(path,{'graduation_year':None},format='json')
        self.assertEqual(AuditEvent.objects.filter(kind='student.update').count(),3)

    def test_student_year_write_requires_current_teaching_scope(self):
        path=f'/api/circles/{self.circle.pk}/members/{self.student.student.pk}'
        for user in [self.student,self.analyst,self.multi]:
            self.assertEqual(self.client_for(user).patch(path,{'graduation_year':2031},format='json').status_code,403)
        self.assertEqual(self.client_for(self.mentor).patch(f'/api/circles/{self.foreign.pk}/members/{self.student.student.pk}',{'graduation_year':2031},format='json').status_code,404)
        self.assertEqual(self.client_for(self.teacher).patch(f'/api/circles/{self.foreign.pk}/members/{self.student.student.pk}',{'graduation_year':2031},format='json').status_code,404)
        s.unenroll(self.teacher,self.circle,self.student.email)
        self.assertEqual(self.client_for(self.teacher).patch(path,{'graduation_year':2031},format='json').status_code,404)
        self.student.student.refresh_from_db();self.assertIsNone(self.student.student.graduation_year)

    def test_student_year_cannot_change_in_archived_circle(self):
        self.circle.archived=True;self.circle.save()
        self.assertEqual(self.client_for(self.admin).patch(f'/api/circles/{self.circle.pk}/members/{self.student.student.pk}',{'graduation_year':2031},format='json').status_code,400)

    def test_student_year_validation_rejects_wrong_types_and_unexpected_fields(self):
        path=f'/api/circles/{self.circle.pk}/members/{self.student.student.pk}'
        client=self.client_for(self.teacher)
        for value in [True,False,2031.5,'2031','',1899,2201,-1,[],{}]:
            with self.subTest(value=value):self.assertEqual(client.patch(path,{'graduation_year':value},format='json').status_code,400)
        for payload in [{},{'graduation_year':2031,'name':'Не изменять'}]:
            self.assertEqual(client.patch(path,payload,format='json').status_code,400)
        self.student.student.refresh_from_db();self.assertIsNone(self.student.student.graduation_year)
        self.assertFalse(AuditEvent.objects.filter(kind='student.update').exists())

    def test_database_rejects_out_of_range_graduation_year(self):
        from django.db import IntegrityError,transaction
        for value in [1899,2201]:
            with self.assertRaises(IntegrityError):
                with transaction.atomic():Student.objects.filter(pk=self.student.student.pk).update(graduation_year=value)


    def _filter_fixture(self):
        pt=[Tag.objects.create(kind='problem',name='Фильтр задачи '+str(i)) for i in range(2)]
        wt=[Tag.objects.create(kind='worksheet',name='Фильтр листка '+str(i)) for i in range(2)]
        topic=Topic.objects.filter(parent__isnull=False,is_other=False,archived=False).first()
        problems=[s.create_material(self.teacher,'problem',{'title':'Фильтр задача '+str(i),'statement':'Условие','solution':'Скрытое решение','topics':[str(topic.pk)],'tags':ids}) for i,ids in enumerate([[pt[0].pk,pt[1].pk],[pt[1].pk],[pt[0].pk]])]
        sheets=[s.create_material(self.teacher,'worksheet',{'title':'Фильтр листок '+str(i),'topics':[str(topic.pk)],'tags':[wt[i].pk],'problems':[str(p.pk) for p in subset]}) for i,subset in enumerate([[problems[0],problems[1]],[problems[0],problems[2]]])]
        for sheet in sheets:s.issue_worksheet(self.teacher,self.circle,sheet.pk,timezone.localdate())
        return pt,wt,problems,sheets,topic

    def test_circle_city_filter_respects_available_circles(self):
        for user,count in [(self.admin,1),(self.teacher,0),(self.analyst,0),(self.student,0)]:
            result=self.client_for(user).get('/api/circles',{'city':self.foreign.city})
            self.assertEqual(result.status_code,200);self.assertEqual(len(result.data),count)
        result=self.client_for(self.teacher).get('/api/circles',{'city':'  '+self.circle.city+'  '})
        self.assertEqual([item['id'] for item in result.data],[str(self.circle.pk)])
        self.assertEqual(self.client_for(self.admin).get('/api/circles',{'city':'Несуществующий город'}).data,[])

    def test_material_tag_filters_combine_with_topic_and_search(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        client=self.client_for(self.teacher)
        response=client.get('/api/materials/problem',{'tag':pt[0].pk,'topic':str(topic.pk),'q':'Фильтр задача 0'})
        self.assertEqual([p['id'] for p in response.data],[str(problems[0].pk)])
        response=client.get('/api/materials/worksheet',{'tag':wt[1].pk})
        self.assertEqual([p['id'] for p in response.data],[str(sheets[1].pk)])
        response=client.get('/api/materials/problem',{'tag':str(pt[0].pk)+','+str(pt[1].pk)})
        self.assertEqual({p['id'] for p in response.data},{str(p.pk) for p in problems})
        self.assertEqual(len(response.data),3)

    def test_material_tag_filter_applies_before_result_limit(self):
        tag=Tag.objects.create(kind='problem',name='Редкий тег')
        topic=Topic.objects.filter(parent__isnull=False,is_other=False,archived=False).first()
        target=s.create_material(self.teacher,'problem',{'title':'Ранняя задача','statement':'Условие','solution':'Решение','topics':[str(topic.pk)],'tags':[tag.pk]})
        Problem.objects.bulk_create([Problem(title='Поздняя задача '+str(i),statement='x',solution='y',author=self.teacher,published=True) for i in range(301)])
        response=self.client_for(self.teacher).get('/api/materials/problem',{'tag':tag.pk})
        self.assertEqual([p['id'] for p in response.data],[str(target.pk)])

    def test_read_filters_reject_malformed_and_wrong_tag_namespaces(self):
        problem_tag=Tag.objects.filter(kind='problem').first();worksheet_tag=Tag.objects.filter(kind='worksheet').first()
        client=self.client_for(self.admin)
        cases=[('/api/materials/problem',{'tag':worksheet_tag.pk}),('/api/materials/worksheet',{'tag':problem_tag.pk}),('/api/analytics',{'problem_tag':worksheet_tag.pk}),('/api/analytics',{'worksheet_tag':problem_tag.pk}),(f'/api/circles/{self.circle.pk}/issues',{'tag':problem_tag.pk})]
        for path,query in cases:self.assertEqual(client.get(path,query).status_code,400)
        for raw in ['abc','0','-1','9999999','1,']:
            self.assertEqual(client.get('/api/materials/problem',{'tag':raw}).status_code,400)

    def test_analytics_intersects_worksheet_and_problem_tags_without_extra_tasks(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        s.mark(self.teacher,self.circle,problems[0].pk,{'student_id':self.student.student.pk,'value':True,'version':0})
        result=self.client_for(self.teacher).get('/api/analytics',{'city':self.circle.city,'worksheet_tag':wt[0].pk,'problem_tag':pt[0].pk,'topic':str(topic.pk)}).data
        self.assertEqual(result['assigned_pairs'],12);self.assertEqual(result['solved_pairs'],1)
        self.assertEqual(result['problem_count'],1);self.assertEqual(result['student_count'],12)
        self.assertEqual(result['rows'][0]['problem_id'],str(problems[0].pk))
        self.assertEqual(result['percent'],8.3)

    def test_analytics_tag_union_deduplicates_repeated_tasks(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        client=self.client_for(self.teacher)
        result=client.get('/api/analytics',{'problem_tag':pt[0].pk}).data
        self.assertEqual(result['assigned_pairs'],24);self.assertEqual(result['problem_count'],2)
        result=client.get('/api/analytics',{'problem_tag':str(pt[0].pk)+','+str(pt[1].pk)}).data
        self.assertEqual(result['assigned_pairs'],36);self.assertEqual(result['problem_count'],3)

    def test_analytics_city_limits_assignments_and_does_not_expand_permissions(self):
        for user in [self.teacher,self.student,self.analyst]:
            result=self.client_for(user).get('/api/analytics',{'city':self.foreign.city})
            self.assertEqual(result.status_code,200);self.assertEqual(result.data['assigned_pairs'],0)
        self.assertEqual(self.client_for(self.teacher).get('/api/analytics',{'city':self.circle.city,'circle':str(self.foreign.pk)}).status_code,404)
        result=self.client_for(self.admin).get('/api/analytics',{'city':self.circle.city,'circle':str(self.foreign.pk)})
        self.assertEqual(result.data['assigned_pairs'],0)

    def test_student_filters_keep_personal_history_after_departure(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        s.unenroll(self.teacher,self.circle,self.student.email);s.unenroll(self.teacher,self.other,self.student.email)
        result=self.client_for(self.student).get('/api/analytics',{'city':self.circle.city,'worksheet_tag':wt[0].pk,'problem_tag':pt[0].pk}).data
        self.assertTrue(result['personal_only']);self.assertEqual(result['assigned_pairs'],1);self.assertEqual(result['student_count'],1)
        self.assertEqual(self.client_for(self.student).get('/api/materials/problem',{'tag':pt[0].pk}).data,[])

    def test_issued_worksheet_filter_includes_tags_and_stays_in_circle(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        result=self.client_for(self.student).get(f'/api/circles/{self.circle.pk}/issues',{'tag':wt[1].pk})
        self.assertEqual(result.status_code,200);self.assertEqual(len(result.data),1)
        self.assertEqual(result.data[0]['worksheet_id'],str(sheets[1].pk))
        self.assertEqual(result.data[0]['tags'],[{'id':wt[1].pk,'public_id':f'TAG-{wt[1].pk}','name':wt[1].name}])
        self.assertEqual(self.client_for(self.student).get(f'/api/circles/{self.foreign.pk}/issues',{'tag':wt[1].pk}).status_code,404)

    def test_archived_tags_remain_available_for_historical_filtering(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        pt[0].archived=True;pt[0].save()
        result=self.client_for(self.teacher).get('/api/materials/problem',{'tag':pt[0].pk})
        self.assertEqual(result.status_code,200);self.assertEqual({p['id'] for p in result.data},{str(problems[i].pk) for i in [0,2]})


    def _report_fixture(self):
        peer=User.objects.get(email='pupil2@plusfolio.example')
        wt=Tag.objects.create(kind='worksheet',name='Срез подробного отчёта')
        pt=Tag.objects.create(kind='problem',name='Одна задача для фильтра')
        ct=Tag.objects.create(kind='circle',name='Область подробного отчёта')
        root=Topic.objects.get(parent=None)
        branch=s.add_topic(self.admin,'Раздел отчёта',root.pk);leaf=s.add_topic(self.admin,'Подраздел отчёта',branch.pk)
        other=s.add_topic(self.admin,'Другой раздел отчёта',root.pk)
        circles=[]
        for i in range(2):
            circle=Circle.objects.create(name='Кружок отчёта '+str(i),city='Город отчёта',created_by=self.admin)
            s.assign_teacher(self.admin,circle,self.teacher.email)
            circles.append(circle)
        circles[0].tags.add(ct);AnalystGrant.objects.create(user=self.analyst,tag=ct)
        for account in [self.student,peer]:s.enroll(self.teacher,circles[0],account.email)
        s.enroll(self.teacher,circles[1],self.student.email)
        problems=[]
        for i,topics in enumerate([[branch.pk,leaf.pk],[leaf.pk],[other.pk]]):
            problems.append(s.create_material(self.teacher,'problem',{'title':'Задача отчёта '+str(i),'statement':'Условие','solution':'Закрытое решение','topics':[str(x) for x in topics],'tags':[pt.pk] if i==0 else []}))
        sheets=[]
        for i,subset in enumerate([[problems[0],problems[1]],[problems[0],problems[2]]]):
            sheet=s.create_material(self.teacher,'worksheet',{'title':'Листок отчёта '+str(i),'topics':[str(branch.pk)],'tags':[wt.pk],'problems':[str(p.pk) for p in subset]})
            sheets.append(sheet);s.issue_worksheet(self.teacher,circles[0],sheet.pk,timezone.localdate())
        s.issue_worksheet(self.teacher,circles[1],sheets[0].pk,timezone.localdate())
        for account,problem in [(self.student,problems[0]),(self.student,problems[1]),(peer,problems[0])]:
            s.mark(self.teacher,circles[0],problem.pk,{'student_id':account.student.pk,'value':True,'version':0})
        for circle,problem in [(circles[0],problems[0]),(circles[0],problems[1]),(circles[1],problems[0]),(circles[1],problems[1])]:
            s.mark(self.teacher,circle,problem.pk,{'value':True,'version':0},discussion=True)
        s.mark(self.teacher,circles[0],problems[0].pk,{'value':True,'version':1},solution_visibility=True)
        return {'peer':peer,'tag':wt,'problem_tag':pt,'circles':circles,'problems':problems,'sheets':sheets,'branch':branch,'leaf':leaf}

    def test_report_deduplicates_tasks_and_complete_worksheets_across_circles(self):
        fixture=self._report_fixture()
        result=self.client_for(self.teacher).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk}).data
        self.assertEqual((result['assigned_pairs'],result['solved_pairs'],result['percent']),(6,3,50))
        self.assertEqual(result['counts']['remaining_pairs'],3)
        self.assertEqual(result['counts']['worksheet_recipient_count'],4)
        self.assertEqual(result['counts']['full_worksheet_count'],1)
        self.assertEqual(result['counts']['issue_count'],3)
        self.assertEqual(result['counts']['worksheet_count'],2)
        self.assertEqual(result['counts']['circle_count'],2)
        self.assertEqual(sum(row['assigned'] for row in result['circle_rows']),8)
        self.assertEqual(sum(row['assigned'] for row in result['worksheet_rows']),8)
        self.assertEqual(result['counts']['students_completed'],0)

    def test_report_full_worksheet_completion_uses_all_tasks_despite_problem_filter(self):
        fixture=self._report_fixture()
        result=self.client_for(self.teacher).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk,'problem_tag':fixture['problem_tag'].pk}).data
        self.assertEqual(result['percent'],100)
        self.assertEqual(result['counts']['full_worksheet_count'],1)
        self.assertEqual(result['counts']['worksheet_recipient_count'],4)
        self.assertTrue(all(row['problem_count']==1 and row['total_problem_count']==2 for row in result['worksheet_rows']))
        self.assertEqual(result['counts']['students_completed'],2)

    def test_report_distribution_and_median_use_each_students_own_denominator(self):
        fixture=self._report_fixture()
        result=self.client_for(self.teacher).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk}).data
        self.assertEqual([row['count'] for row in result['distribution']],[0,0,1,1,0,0])
        self.assertEqual(result['counts']['median_percent'],50)
        self.assertEqual(sorted(row['percent'] for row in result['student_rows']),[33.3,66.7])
        fresh=User.objects.get(email='pupil3@plusfolio.example')
        s.enroll(self.teacher,fixture['circles'][0],fresh.email)
        result=self.client_for(self.teacher).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk}).data
        self.assertEqual(result['counts']['students_without_solves'],1)
        self.assertEqual(result['distribution'][0]['count'],1)
        self.assertEqual(sum(row['count'] for row in result['distribution']),3)
        self.assertEqual(result['counts']['median_percent'],33.3)

    def test_report_topic_ancestors_deduplicate_overlapping_topic_links(self):
        fixture=self._report_fixture()
        result=self.client_for(self.teacher).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk}).data
        for node,depth in [(fixture['branch'],1),(fixture['leaf'],2)]:
            row=next(row for row in result['topic_rows'] if row['id']==str(node.pk))
            self.assertEqual((row['assigned'],row['solved'],row['percent']),(4,3,75))
            self.assertEqual((row['recipients'],row['problem_count'],row['depth']),(2,2,depth))
        self.assertEqual(sum(row['assigned'] for row in result['topic_rows'] if row['depth']==1),6)

    def test_report_discussion_is_local_and_pending_counts_unique_tasks(self):
        fixture=self._report_fixture()
        result=self.client_for(self.teacher).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk}).data
        self.assertEqual(result['discussion'],{'total':5,'discussed':4,'remaining':1,'visible':1,'percent':80})
        self.assertEqual(result['attention']['discussed_pending_problem_count'],1)
        self.assertEqual(result['attention']['unsolved_problem_count'],1)
        self.assertEqual(result['attention']['lowest_closure'][0]['problem_id'],str(fixture['problems'][2].pk))

    def test_report_personal_response_has_no_peer_lists_or_group_statistics(self):
        fixture=self._report_fixture()
        for circle in fixture['circles']:s.unenroll(self.teacher,circle,self.student.email)
        response=self.client_for(self.student).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk})
        result=response.data
        self.assertTrue(result['personal_only'])
        self.assertEqual((result['assigned_pairs'],result['solved_pairs']),(3,2))
        self.assertEqual(result['counts']['full_worksheet_count'],1)
        self.assertEqual(result['counts']['worksheet_recipient_count'],2)
        for key in ['student_rows','circle_rows','distribution']:self.assertEqual(result[key],[])
        self.assertIsNone(result['discussion']);self.assertIsNone(result['attention']['discussed_pending_problem_count'])
        self.assertNotIn(fixture['peer'].name,response.content.decode())
        self.assertTrue(all(row['recipients']==1 for row in result['worksheet_rows']))

    def test_report_analyst_scope_filters_every_new_breakdown(self):
        fixture=self._report_fixture();hidden=User.objects.get(email='pupil3@plusfolio.example')
        s.enroll(self.teacher,fixture['circles'][1],hidden.email)
        response=self.client_for(self.analyst).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk})
        result=response.data
        self.assertEqual({row['id'] for row in result['circle_rows']},{str(fixture['circles'][0].pk)})
        self.assertEqual({row['id'] for row in result['student_rows']},{self.student.student.pk,fixture['peer'].student.pk})
        self.assertEqual(result['discussion']['total'],3)
        self.assertNotIn(hidden.name,response.content.decode())
        self.assertNotIn(fixture['circles'][1].name,response.content.decode())
        self.assertTrue(all(row['circle_count']==1 for row in result['worksheet_rows']))

    def test_report_mixed_roles_do_not_expose_peers_from_student_only_circle(self):
        fixture=self._report_fixture()
        self.student.roles.append('teacher');self.student.save();s.profiles(self.student)
        s.assign_teacher(self.admin,fixture['circles'][1],self.student.email)
        result=self.client_for(self.student).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk}).data
        self.assertFalse(result['personal_only'])
        self.assertEqual({row['id'] for row in result['student_rows']},{self.student.student.pk})
        self.assertEqual({row['id'] for row in result['circle_rows']},{str(fixture['circles'][1].pk)})
        self.assertEqual(result['discussion']['total'],2)

    def test_report_empty_slice_has_no_undefined_ratios(self):
        tag=Tag.objects.create(kind='worksheet',name='Нет назначений')
        response=self.client_for(self.admin).get('/api/analytics',{'worksheet_tag':tag.pk})
        result=response.data
        self.assertEqual(result['assigned_pairs'],0);self.assertIsNone(result['percent']);self.assertIsNone(result['counts']['median_percent'])
        for key in ['rows','worksheet_rows','topic_rows','student_rows','circle_rows']:self.assertEqual(result[key],[])
        self.assertEqual(result['counts']['worksheet_recipient_count'],0)
        self.assertEqual(result['discussion']['total'],0);self.assertIsNone(result['discussion']['percent'])
        self.assertNotIn('NaN',response.content.decode())

    def test_report_grades_and_queries_do_not_require_per_student_lookups(self):
        from datetime import date
        from unittest.mock import patch
        from django.test.utils import CaptureQueriesContext
        from django.db import connection
        fixture=self._report_fixture()
        Student.objects.filter(pk=self.student.student.pk).update(graduation_year=2031)
        client=self.client_for(self.teacher)
        with patch('core.school.timezone.localdate',return_value=date(2026,10,21)),CaptureQueriesContext(connection) as queries:
            result=client.get('/api/analytics',{'worksheet_tag':fixture['tag'].pk}).data
        self.assertLessEqual(len(queries),30)
        row=next(row for row in result['student_rows'] if row['id']==self.student.student.pk)
        self.assertEqual(row['grade_label'],'7 класс')


    def test_admin_creates_tags_for_three_separate_dictionaries_with_audit(self):
        client=self.client_for(self.admin)
        for kind in ['problem','worksheet','circle']:
            response=client.post('/api/tags',{'kind':kind,'name':'  Новый тег  '},format='json')
            self.assertEqual(response.status_code,201)
            tag=Tag.objects.get(kind=kind,name='Новый тег')
            self.assertFalse(tag.archived)
            self.assertIn({'id':tag.pk,'public_id':f'TAG-{tag.pk}','kind':kind,'name':'Новый тег','archived':False},response.data)
            event=AuditEvent.objects.get(kind='tag.create',payload__id=tag.pk)
            self.assertEqual(event.actor,self.admin);self.assertIsNone(event.circle)
            self.assertEqual(event.payload,{'id':tag.pk,'kind':kind,'name':'Новый тег'})

    def test_tag_creation_forbidden_for_every_non_admin_role(self):
        before=Tag.objects.count()
        for user in [self.teacher,self.mentor,self.student,self.analyst,self.multi]:
            result=self.client_for(user).post('/api/tags',{'kind':'problem','name':'Нельзя'},format='json')
            self.assertEqual(result.status_code,403)
        self.assertEqual(Tag.objects.count(),before)
        self.assertFalse(AuditEvent.objects.filter(kind='tag.create').exists())

    def test_tag_creation_rejects_invalid_dictionary_name_type_and_length(self):
        client=self.client_for(self.admin);before=Tag.objects.count()
        payloads=[{}, {'kind':'topic','name':'Тег'}, {'kind':None,'name':'Тег'}]
        payloads.extend({'kind':'problem','name':name} for name in ['', '   ',None,True,123,[],{},'x'*101])
        for payload in payloads:
            with self.subTest(payload=payload):self.assertEqual(client.post('/api/tags',payload,format='json').status_code,400)
        self.assertEqual(Tag.objects.count(),before)
        self.assertFalse(AuditEvent.objects.filter(kind='tag.create').exists())

    def test_tag_duplicate_is_rejected_only_within_its_dictionary(self):
        client=self.client_for(self.admin)
        for kind in ['problem','worksheet']:
            self.assertEqual(client.post('/api/tags',{'kind':kind,'name':'Одинаковое название'},format='json').status_code,201)
        response=client.post('/api/tags',{'kind':'problem','name':' Одинаковое название '},format='json')
        self.assertEqual(response.status_code,409)
        self.assertIn('уже есть',response.data['detail'])
        self.assertEqual(Tag.objects.filter(name='Одинаковое название').count(),2)
        self.assertEqual(AuditEvent.objects.filter(kind='tag.create').count(),2)

    def test_new_admin_tag_is_available_to_teacher_for_publishing(self):
        self.client_for(self.admin).post('/api/tags',{'kind':'problem','name':'Новый доступный тег'},format='json')
        tag=Tag.objects.get(kind='problem',name='Новый доступный тег')
        teacher=self.client_for(self.teacher)
        self.assertIn(tag.pk,[row['id'] for row in teacher.get('/api/tags').data])
        topic=Topic.objects.filter(parent__isnull=False,is_other=False,archived=False).first()
        result=teacher.post('/api/materials/problem',{'title':'Задача с новым тегом','statement':'Условие','solution':'Решение','topics':[str(topic.pk)],'tags':[tag.pk]},format='json')
        self.assertEqual(result.status_code,201)
        self.assertEqual(result.data['tags'],[{'id':tag.pk,'public_id':f'TAG-{tag.pk}','name':tag.name}])

    def test_tag_reading_remains_available_for_filters_of_all_roles(self):
        for user in [self.admin,self.teacher,self.student,self.analyst,self.multi]:
            response=self.client_for(user).get('/api/tags')
            self.assertEqual(response.status_code,200)
            self.assertEqual(len(response.data),Tag.objects.count())


    def test_analytics_topic_tree_unites_multiple_branches_with_descendants(self):
        fixture=self._report_fixture()
        other=fixture['problems'][2].topics.first()
        params={'worksheet_tag':fixture['tag'].pk,'topic':str(fixture['branch'].pk)+','+str(other.pk)}
        result=self.client_for(self.teacher).get('/api/analytics',params).data
        self.assertEqual((result['assigned_pairs'],result['solved_pairs'],result['problem_count']),(6,3,3))
        self.assertEqual({row['problem_id'] for row in result['rows']},{str(problem.pk) for problem in fixture['problems']})
        self.assertEqual(result['counts']['full_worksheet_count'],1)

    def test_analytics_topic_tree_overlapping_parent_child_selections_do_not_duplicate(self):
        fixture=self._report_fixture()
        branch=str(fixture['branch'].pk);leaf=str(fixture['leaf'].pk)
        result=self.client_for(self.teacher).get('/api/analytics',{'worksheet_tag':fixture['tag'].pk,'topic':','.join([branch,leaf,leaf])}).data
        self.assertEqual((result['assigned_pairs'],result['solved_pairs'],result['problem_count']),(4,3,2))
        self.assertEqual(result['percent'],75)
        self.assertEqual(result['student_count'],2)

    def test_analytics_topic_tree_preserves_role_scope_for_multiple_branches(self):
        fixture=self._report_fixture();other=fixture['problems'][2].topics.first()
        params={'worksheet_tag':fixture['tag'].pk,'topic':str(fixture['branch'].pk)+','+str(other.pk)}
        student=self.client_for(self.student).get('/api/analytics',params).data
        self.assertTrue(student['personal_only']);self.assertEqual(student['assigned_pairs'],3)
        self.assertEqual(student['student_rows'],[]);self.assertEqual(student['circle_rows'],[])
        analyst=self.client_for(self.analyst).get('/api/analytics',params).data
        self.assertEqual(analyst['assigned_pairs'],6)
        self.assertEqual({row['id'] for row in analyst['circle_rows']},{str(fixture['circles'][0].pk)})
        params['topic']+=',00000000-0000-0000-0000-000000000000'
        self.assertEqual(self.client_for(self.teacher).get('/api/analytics',params).status_code,400)


    def test_multiple_cities_use_union_after_permission_scoping(self):
        cities=[self.circle.city,self.foreign.city,self.circle.city]
        for user,expected in [(self.admin,{self.circle.pk,self.foreign.pk}),(self.teacher,{self.circle.pk}),(self.student,{self.circle.pk})]:
            response=self.client_for(user).get('/api/circles',{'city':cities})
            self.assertEqual(response.status_code,200)
            self.assertEqual({row['id'] for row in response.data},{str(pk) for pk in expected})
        # Repeated parameters preserve a comma inside a real city name.
        circle=Circle.objects.create(name='Проверка названия города',city='Город, район',created_by=self.admin)
        response=self.client_for(self.admin).get('/api/circles',{'city':['Город, район',self.circle.city]})
        self.assertEqual({row['id'] for row in response.data},{str(circle.pk),str(self.circle.pk)})

    def test_multiple_circles_and_cities_intersect_without_duplicate_results(self):
        fixture=self._report_fixture();first,second=fixture['circles']
        second.city='Ещё один город';second.save()
        params={'circle':[str(first.pk),str(second.pk),str(first.pk)],'city':[first.city,second.city],'worksheet_tag':fixture['tag'].pk}
        client=self.client_for(self.teacher)
        result=client.get('/api/analytics',params)
        self.assertEqual(result.status_code,200)
        self.assertEqual((result.data['assigned_pairs'],result.data['solved_pairs']),(6,3))
        self.assertEqual(result.data['counts']['circle_count'],2)
        params['city']=[second.city]
        result=client.get('/api/analytics',params).data
        self.assertEqual((result['assigned_pairs'],result['solved_pairs']),(2,2))
        self.assertEqual([row['id'] for row in result['circle_rows']],[str(second.pk)])

    def test_multiple_circle_filter_rejects_any_unavailable_or_invalid_circle(self):
        client=self.client_for(self.teacher)
        for raw,status in [([str(self.circle.pk),str(self.foreign.pk)],404),([str(self.circle.pk),'bad-id'],400),([str(self.circle.pk),''],400)]:
            response=client.get('/api/analytics',{'circle':raw,'city':self.circle.city})
            self.assertEqual(response.status_code,status)
        response=client.get('/api/analytics',{'circle':str(self.circle.pk)+','+str(self.circle.pk)})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.data['counts']['circle_count'],1)

    def test_repeated_tag_and_topic_filters_union_and_preserve_namespace_checks(self):
        pt,wt,problems,sheets,topic=self._filter_fixture();client=self.client_for(self.teacher)
        params={'tag':[str(pt[0].pk),str(pt[1].pk)],'topic':[str(topic.pk),str(topic.pk)]}
        result=client.get('/api/materials/problem',params)
        self.assertEqual({row['id'] for row in result.data},{str(p.pk) for p in problems})
        self.assertEqual(len(result.data),3)
        result=client.get('/api/analytics',{'city':[self.circle.city,self.foreign.city],'problem_tag':[p.pk for p in pt],'worksheet_tag':[w.pk for w in wt]})
        self.assertEqual((result.data['assigned_pairs'],result.data['problem_count']),(36,3))
        result=client.get(f'/api/circles/{self.circle.pk}/issues',{'tag':[t.pk for t in wt]})
        self.assertEqual({row['worksheet_id'] for row in result.data},{str(sheet.pk) for sheet in sheets})
        params['tag'].append(str(wt[0].pk))
        self.assertEqual(client.get('/api/materials/problem',params).status_code,400)

    def test_multiple_city_reports_preserve_personal_history_and_hide_peer_analytics(self):
        fixture=self._report_fixture();first,second=fixture['circles']
        second.city='Второй город ученика';second.save()
        client=self.client_for(self.student)
        params={'city':[first.city,second.city],'circle':[str(first.pk),str(second.pk)],'worksheet_tag':fixture['tag'].pk}
        result=client.get('/api/analytics',params).data
        self.assertTrue(result['personal_only']);self.assertEqual(result['assigned_pairs'],3)
        self.assertEqual(result['student_rows'],[]);self.assertEqual(result['circle_rows'],[])
        for circle in fixture['circles']:s.unenroll(self.teacher,circle,self.student.email)
        params.pop('circle')
        result=client.get('/api/analytics',params).data
        self.assertEqual(result['assigned_pairs'],3);self.assertEqual(result['student_count'],1)

    def test_batch_worksheet_issue_deduplicates_and_assigns_every_selected_sheet(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        circle=Circle.objects.create(name='Пакетная выдача',created_by=self.admin)
        s.assign_teacher(self.admin,circle,self.teacher.email);s.enroll(self.teacher,circle,self.student.email)
        response=self.client_for(self.teacher).post(f'/api/circles/{circle.pk}/issues',{'worksheet_ids':[str(sheets[0].pk),str(sheets[1].pk),str(sheets[0].pk)],'issued_on':timezone.localdate().isoformat()},format='json')
        self.assertEqual(response.status_code,200)
        self.assertEqual(circle.issues.count(),2)
        self.assertEqual(Assignment.objects.filter(issue__circle=circle,student=self.student.student).count(),2)

    def test_batch_worksheet_issue_rolls_back_all_sheets_assignments_and_audit(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        circle=Circle.objects.create(name='Атомарная выдача',created_by=self.admin)
        s.assign_teacher(self.admin,circle,self.teacher.email);s.enroll(self.teacher,circle,self.student.email)
        sheets[1].archived=True;sheets[1].save()
        response=self.client_for(self.teacher).post(f'/api/circles/{circle.pk}/issues',{'worksheet_ids':[str(sheet.pk) for sheet in sheets],'issued_on':timezone.localdate().isoformat()},format='json')
        self.assertEqual(response.status_code,400)
        self.assertFalse(circle.issues.exists());self.assertFalse(Assignment.objects.filter(issue__circle=circle).exists())
        self.assertFalse(AuditEvent.objects.filter(circle=circle,kind='worksheet.issue').exists())

    def test_empty_or_malformed_batch_is_rejected_without_writes(self):
        client=self.client_for(self.admin);before=WorksheetIssue.objects.count();invites=Invitation.objects.count()
        for values in [[],str(self.issue.worksheet_id),[None],[''],['bad-id'],[str(self.issue.worksheet_id)+','+str(self.issue.worksheet_id)]]:
            with self.subTest(values=values):
                self.assertEqual(client.post(f'/api/circles/{self.circle.pk}/issues',{'worksheet_ids':values,'issued_on':timezone.localdate().isoformat()},format='json').status_code,400)
                self.assertEqual(client.post('/api/invitations',{'circle_ids':values,'role':'student','email':'batch@plusfolio.example'},format='json').status_code,400)
        self.assertEqual(WorksheetIssue.objects.count(),before);self.assertEqual(Invitation.objects.count(),invites)

    def test_batch_invitation_creates_separate_scoped_letters_without_duplicates(self):
        response=self.client_for(self.admin).post('/api/invitations',{'circle_ids':[str(self.circle.pk),str(self.other.pk),str(self.circle.pk)],'role':'student','email':'batch@plusfolio.example'},format='json')
        self.assertEqual(response.status_code,201);self.assertEqual(len(response.data['ids']),2)
        objects=Invitation.objects.filter(pk__in=response.data['ids'])
        self.assertEqual(set(objects.values_list('circle_id',flat=True)),{self.circle.pk,self.other.pk})
        self.assertEqual(EmailOutbox.objects.filter(invitation__in=objects).count(),2)
        self.assertTrue(all(row.role=='student' for row in objects))

    def test_batch_invitation_failure_rolls_back_revocations_and_outbox(self):
        previous=s.invite(self.admin,'batch@plusfolio.example','student',self.circle)
        self.other.archived=True;self.other.save()
        counts=(Invitation.objects.count(),EmailOutbox.objects.count(),AuditEvent.objects.count())
        response=self.client_for(self.admin).post('/api/invitations',{'circle_ids':[str(self.circle.pk),str(self.other.pk)],'role':'student','email':'batch@plusfolio.example'},format='json')
        self.assertEqual(response.status_code,403)
        previous.refresh_from_db();self.assertEqual(previous.state,'pending')
        self.assertEqual((Invitation.objects.count(),EmailOutbox.objects.count(),AuditEvent.objects.count()),counts)
        self.assertEqual(self.client_for(self.teacher).post('/api/invitations',{'circle_ids':[str(self.circle.pk),str(self.foreign.pk)],'role':'student','email':'batch@plusfolio.example'},format='json').status_code,404)


    def test_account_returns_own_identity_and_only_assigned_role_sections(self):
        for user in [self.admin,self.teacher,self.student,self.analyst,self.multi]:
            result=self.client_for(user).get('/api/account',{'user_id':self.admin.pk})
            self.assertEqual(result.status_code,200)
            self.assertEqual(result.data['user']['id'],user.pk)
            for role in ['student','teacher','analyst','admin']:
                self.assertEqual(result.data[role] is not None,role in user.roles)
            self.assertNotIn('password',result.data['user'])
            self.assertTrue(all(set(circle)<= {'id','public_id','name','city','archived','relations'} for circle in result.data['circles']))

    def test_account_circle_overview_cannot_expand_role_scope(self):
        for user in [self.student,self.teacher,self.analyst,self.multi]:
            result=self.client_for(user).get('/api/account').data
            expected={str(circle.pk) for circle in s.circles_for(user)}
            self.assertEqual({circle['id'] for circle in result['circles']},expected)
            self.assertNotIn(str(self.foreign.pk),expected)
        result=self.client_for(self.admin).get('/api/account').data
        self.assertEqual(result['admin']['active_users'],17)
        self.assertEqual(result['admin']['active_circles'],3)

    def test_account_personal_progress_is_unique_and_survives_departure_and_role_union(self):
        client=self.client_for(self.student);before=client.get('/api/account').data['student']
        fixture=self._report_fixture()
        # Two sheets and two circles share tasks; peers' solves never enter the personal numerator.
        result=client.get('/api/account').data['student']
        self.assertEqual(result['assigned'],before['assigned']+3)
        self.assertEqual(result['solved'],before['solved']+2)
        for circle in fixture['circles']:s.unenroll(self.teacher,circle,self.student.email)
        self.student.roles.append('teacher');self.student.save();s.profiles(self.student)
        s.assign_teacher(self.admin,self.circle,self.student.email)
        combined=client.get('/api/account').data
        self.assertEqual(combined['student']['assigned'],result['assigned'])
        self.assertEqual(combined['student']['solved'],result['solved'])
        self.assertIsNotNone(combined['teacher'])
        self.assertEqual(combined['student']['active_circles'],before['active_circles'])

    def test_account_teacher_summary_excludes_ended_and_archived_teaching(self):
        baseline=self.client_for(self.teacher).get('/api/account').data['teacher']
        fixture=self._report_fixture();first,second=fixture['circles']
        first.archived=True;first.save()
        s.assign_teacher(self.admin,second,self.teacher.email,remove=True)
        before=self.client_for(self.teacher).get('/api/account').data['teacher']
        self.assertEqual(before['active_circles'],baseline['active_circles'])
        self.assertEqual(before['students'],baseline['students'])
        # Other role grants do not become teaching assignments.
        self.teacher.roles.append('analyst');self.teacher.save()
        tag=Tag.objects.create(kind='circle',name='Только аналитика кабинета');self.foreign.tags.add(tag)
        AnalystGrant.objects.create(user=self.teacher,tag=tag)
        result=self.client_for(self.teacher).get('/api/account').data
        self.assertEqual(result['teacher'],before)
        self.assertEqual(result['analyst']['active_circles'],1)
        self.assertEqual(result['analyst']['tags'],[{'tag__id':tag.pk,'tag__name':tag.name}])

    def test_account_name_update_is_self_only_and_audited(self):
        original_admin=self.admin.name
        response=self.client_for(self.student).patch('/api/account',{'name':'  Иванов   Иван Иванович  ','graduation_year':2031,'no_patronymic':False},format='json')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.data['user']['name'],'Иванов Иван Иванович')
        self.student.refresh_from_db();self.admin.refresh_from_db()
        self.assertEqual(self.student.name,'Иванов Иван Иванович');self.assertEqual(self.admin.name,original_admin)
        self.assertEqual(AuditEvent.objects.get(kind='account.student_profile').actor_id,self.student.pk)
        self.assertEqual(response.data['user']['roles'],['student'])

    def test_account_profile_validation_rejects_role_email_and_education_changes(self):
        client=self.client_for(self.student)
        for payload in [{},{'name':''},{'name':'   '},{'name':'x'*121},{'name':None},{'name':[]},{'name':123},{'name':'Имя','roles':['admin']},{'email':'other@plusfolio.example'},{'graduation_year':2031},{'id':self.admin.pk,'name':'Другое имя'}]:
            self.assertEqual(client.patch('/api/account',payload,format='json').status_code,400)
        self.student.refresh_from_db();self.assertEqual(self.student.roles,['student'])
        self.assertFalse(AuditEvent.objects.filter(kind='account.name').exists())

    def test_account_noop_name_update_does_not_create_audit_event(self):
        self.assertEqual(self.client_for(self.teacher).patch('/api/account',{'name':self.teacher.name},format='json').status_code,200)
        self.assertFalse(AuditEvent.objects.filter(kind='account.name').exists())

    def test_account_anonymous_and_csrf_writes_are_rejected(self):
        anonymous=APIClient()
        self.assertEqual(anonymous.get('/api/account').status_code,403)
        self.assertEqual(anonymous.post('/api/account/password',{},format='json').status_code,403)
        client=APIClient(enforce_csrf_checks=True);client.force_login(self.student)
        self.assertEqual(client.patch('/api/account',{'name':'Без CSRF'},format='json').status_code,403)
        self.assertEqual(client.post('/api/account/password',{},format='json').status_code,403)

    def test_account_password_validates_current_password_confirmation_and_input_types(self):
        client=self.client_for(self.student);new='AccountTestingPass2026!'
        good={'current_password':'PlusfolioDemo2026!','new_password':new,'confirm_password':new}
        for payload in [{},{**good,'current_password':'wrong'},{**good,'confirm_password':'other'},{**good,'new_password':None},{**good,'roles':['admin']}]:
            self.assertEqual(client.post('/api/account/password',payload,format='json').status_code,400)
        self.student.refresh_from_db();self.assertTrue(self.student.check_password('PlusfolioDemo2026!'))
        self.assertFalse(AuditEvent.objects.filter(kind='account.password').exists())

    @override_settings(AUTH_PASSWORD_VALIDATORS=[{'NAME':'django.contrib.auth.password_validation.MinimumLengthValidator','OPTIONS':{'min_length':12}}])
    def test_account_password_runs_django_password_validators(self):
        response=self.client_for(self.student).post('/api/account/password',{'current_password':'PlusfolioDemo2026!','new_password':'short','confirm_password':'short'},format='json')
        self.assertEqual(response.status_code,400)
        self.student.refresh_from_db();self.assertTrue(self.student.check_password('PlusfolioDemo2026!'))

    def test_account_password_preserves_current_session_revokes_others_and_logs_no_secrets(self):
        current=self.client_for(self.student);other=self.client_for(self.student);new='AccountTestingPass2026!'
        response=current.post('/api/account/password',{'current_password':'PlusfolioDemo2026!','new_password':new,'confirm_password':new},format='json')
        self.assertEqual(response.status_code,200)
        self.assertEqual(current.get('/api/session').json()['user']['id'],self.student.pk)
        self.assertIsNone(other.get('/api/session').json()['user'])
        self.student.refresh_from_db();self.assertTrue(self.student.check_password(new))
        self.assertFalse(self.student.check_password('PlusfolioDemo2026!'))
        event=AuditEvent.objects.get(kind='account.password')
        self.assertEqual(event.payload,{})
        self.assertNotIn(new,str(response.data))


    def test_my_tasks_only_include_received_tasks_and_deduplicate_shared_sheets(self):
        client=self.client_for(self.student)
        expected=set(Problem.objects.filter(items__worksheet__issues__assignments__student=self.student.student).values_list('pk',flat=True))
        result=client.get('/api/my-materials/problem')
        self.assertEqual(result.status_code,200);self.assertEqual(len(result.data),len(expected))
        self.assertEqual({row['id'] for row in result.data},{str(pk) for pk in expected})
        hidden=Problem.objects.filter(items__worksheet__issues__circle=self.foreign).exclude(pk__in=expected).first()
        if hidden:self.assertEqual(client.get('/api/my-problems/'+str(hidden.pk)).status_code,404)
        self.assertEqual(self.client_for(self.teacher).get('/api/my-materials/problem').status_code,403)

    def test_my_worksheets_keep_every_issue_with_its_circle_date_and_number(self):
        result=self.client_for(self.student).get('/api/my-materials/worksheet')
        expected=WorksheetIssue.objects.filter(assignments__student=self.student.student).distinct()
        self.assertEqual(len(result.data),expected.count())
        self.assertEqual({row['id'] for row in result.data},set(expected.values_list('pk',flat=True)))
        self.assertTrue(all(row['number']>=1 and row['pdf_url'].startswith('/api/my-issues/') for row in result.data))
        self.assertEqual(self.client_for(self.student).get('/api/my-issues/'+str(self.foreign.issues.first().pk)).status_code,404)

    def test_my_task_solution_is_global_across_current_circles_and_live(self):
        client=self.client_for(self.student)
        # This seeded task belongs to both circles; expose it in the first circle only.
        discussion=s.mark(self.teacher,self.circle,self.problem.pk,{'value':True,'version':Discussion.objects.filter(circle=self.circle,problem=self.problem).values_list('version',flat=True).first() or 0},discussion=True)
        s.mark(self.teacher,self.circle,self.problem.pk,{'value':True,'version':discussion['version']},solution_visibility=True)
        own=client.get('/api/my-problems/'+str(self.problem.pk))
        self.assertEqual(own.data['solution'],self.problem.solution)
        circle_detail=client.get('/api/materials/problem/'+str(self.problem.pk),{'circle':str(self.other.pk)})
        self.assertNotIn('solution',circle_detail.data)
        count=len(client.get('/api/my-materials/problem').data)
        s.unenroll(self.teacher,self.circle,self.student.email)
        self.assertNotIn('solution',client.get('/api/my-problems/'+str(self.problem.pk)).data)
        s.unenroll(self.teacher,self.other,self.student.email)
        self.assertEqual(len(client.get('/api/my-materials/problem').data),count)
        self.assertEqual(client.get('/api/my-problems/'+str(self.problem.pk)).status_code,200)

    def test_issue_numbers_are_unique_and_do_not_change_for_backdated_new_issues(self):
        pt,wt,problems,sheets,topic=self._filter_fixture()
        numbers=list(self.circle.issues.order_by('number').values_list('number',flat=True))
        self.assertEqual(numbers,list(range(1,len(numbers)+1)))
        before={issue.pk:issue.number for issue in self.circle.issues.all()}
        extra=s.create_material(self.teacher,'worksheet',{'title':'Новая выдача задним числом','topics':[str(topic.pk)],'tags':[],'problems':[str(problems[0].pk)]})
        issue=s.issue_worksheet(self.teacher,self.circle,extra.pk,timezone.localdate()-timedelta(days=20))
        self.assertEqual(issue.number,len(numbers)+1)
        self.assertEqual({issue.pk:issue.number for issue in self.circle.issues.filter(pk__in=before)},before)

    def test_tex_projects_have_separate_parameters_task_files_and_filtered_solutions(self):
        import io,zipfile
        path='/api/materials/worksheet/'+str(self.issue.worksheet_id)
        response=self.client_for(self.teacher).get(path+'/project',{'solutions':'1'})
        self.assertEqual(response.status_code,200)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            names=archive.namelist();self.assertTrue({'main.tex','variables.tex','template.tex','problems.tex','solutions.tex'}<=set(names))
            self.assertEqual(len([name for name in names if name.startswith('tasks/')]),self.issue.worksheet.items.count())
            self.assertEqual(len([name for name in names if name.startswith('solutions/')]),self.issue.worksheet.items.count())
            self.assertIn('НАЗВАНИЕ КРУЖКА',archive.read('variables.tex').decode())
            self.assertIn('tikz',archive.read('template.tex').decode())
        self.assertEqual(self.client_for(self.student).get(path+'/project',{'solutions':'1'}).status_code,403)

    def test_flattened_tex_is_standalone_and_context_replaces_placeholders(self):
        response=self.client_for(self.teacher).get(f'/api/circles/{self.circle.pk}/issues/{self.issue.pk}/tex',{'solutions':'1'})
        self.assertEqual(response.status_code,200);text=response.content.decode()
        self.assertNotIn(r'\input{tasks/',text);self.assertNotIn(r'\input{variables.tex}',text)
        self.assertIn(self.circle.name,text);self.assertIn(self.issue.issued_on.strftime('%d.%m.%Y'),text)
        self.assertIn(r'\newcommand{\PlusfolioNumber}{'+str(self.issue.number)+'}',text)
        self.assertNotIn('ДАТА ВЫДАЧИ',text)

    def test_pdf_project_solution_mask_is_rechecked_before_using_cache(self):
        from unittest.mock import patch
        captured=[]
        def compile(project):captured.append(project);return b'%PDF-1.4 test'
        client=self.client_for(self.student);path=f'/api/my-issues/{self.issue.pk}/pdf'
        with patch('core.documents.render_pdf',side_effect=compile):
            self.assertEqual(client.get(path,{'solutions':'1'}).status_code,200)
            self.assertFalse(any(name.startswith('solutions/') for name in captured[-1]))
            discussed=s.mark(self.teacher,self.circle,self.problem.pk,{'value':True,'version':Discussion.objects.filter(circle=self.circle,problem=self.problem).values_list('version',flat=True).first() or 0},discussion=True)
            shown=s.mark(self.teacher,self.circle,self.problem.pk,{'value':True,'version':discussed['version']},solution_visibility=True)
            self.assertEqual(client.get(path,{'solutions':'1'}).status_code,200)
            self.assertEqual(len([name for name in captured[-1] if name.startswith('solutions/')]),1)
            s.mark(self.teacher,self.circle,self.problem.pk,{'value':False,'version':shown['version']},solution_visibility=True)
            self.assertEqual(client.get(path,{'solutions':'1'}).status_code,200)
            self.assertFalse(any(name.startswith('solutions/') for name in captured[-1]))

    def test_foreign_pdf_and_tex_are_unavailable_and_personal_pdf_survives_departure(self):
        from unittest.mock import patch
        client=self.client_for(self.student)
        with patch('core.documents.render_pdf',return_value=b'%PDF-1.4 test'):
            self.assertEqual(client.get(f'/api/circles/{self.foreign.pk}/issues/{self.foreign.issues.first().pk}/pdf').status_code,404)
            self.assertEqual(client.get(f'/api/materials/worksheet/{self.foreign.issues.first().worksheet_id}/pdf').status_code,404)
            s.unenroll(self.teacher,self.circle,self.student.email)
            self.assertEqual(client.get(f'/api/my-issues/{self.issue.pk}/pdf').status_code,200)
            self.assertEqual(client.get(f'/api/circles/{self.circle.pk}/issues/{self.issue.pk}/pdf').status_code,404)

    def test_plain_fallback_escapes_tex_and_optional_sources_are_immutable(self):
        from .documents import fragment
        from types import SimpleNamespace
        item=SimpleNamespace(statement='100% #1 & 2_3 $x^2$ '+r'\input{secret}',statement_tex='')
        text=fragment(item,'statement')
        self.assertIn(r'100\% \#1 \& 2\_3 $x^2$',text)
        self.assertIn(r'\textbackslash{}input\{secret\}',text)
        problem=self.problem;problem.statement_tex=r'\begin{tikzpicture}\draw (0,0)--(1,1);\end{tikzpicture}'
        with self.assertRaises(ValidationError):problem.save()

    def test_invalid_tex_worksheet_does_not_publish_or_leave_items(self):
        from unittest.mock import patch
        topic=Topic.objects.filter(is_other=True,archived=False).first()
        problem=s.create_material(self.teacher,'problem',{'title':'TeX only','topics':[str(topic.pk)],'statement_tex':r'\unknowncommand','solution':'Ответ'})
        self.assertEqual(problem.statement,'')
        before=(Worksheet.objects.count(),WorksheetItem.objects.count())
        with patch('core.documents.render_pdf',side_effect=ValidationError('Broken TeX')):
            with self.assertRaises(ValidationError):s.create_material(self.teacher,'worksheet',{'title':'Broken sheet','topics':[str(topic.pk)],'problems':[str(problem.pk)]})
        self.assertEqual((Worksheet.objects.count(),WorksheetItem.objects.count()),before)



    def _invite_token(self,invitation):
        import json
        return json.loads(s.cipher().decrypt(invitation.letters.get().encrypted_payload.encode()))['token']

    def _named_student(self,email='semenov@plusfolio.example',name='Семёнов Алексей Андреевич',year=2031):
        user=User.objects.create_user(email,'unused',name=name,roles=['student']);s.profiles(user)
        user.student.graduation_year=year;user.student.save()
        return user

    def test_people_admin_filters_are_private_composable_and_unicode_aware(self):
        pupil=self._named_student();pupil.roles.append('teacher');pupil.save();s.profiles(pupil)
        inactive=self._named_student('inactive@plusfolio.example','Семенов Алексей Петрович',2032)
        inactive.is_active=False;inactive.save()
        client=self.client_for(self.admin)
        result=client.get('/api/users',{'meta':'1','q':'АЛЕКСЕЙ семенов','role':['student','analyst'],'state':'active','graduation_year':['2031','2032']})
        self.assertEqual(result.status_code,200)
        self.assertEqual([row['id'] for row in result.data['results']],[pupil.pk])
        self.assertEqual(result.data['count'],1)
        self.assertEqual(client.get('/api/users',{'meta':'1','q':'inactive@','state':'inactive'}).data['count'],1)
        self.assertEqual(client.get('/api/users',{'meta':'1','role':['teacher','analyst']}).data['count'],5)
        for other in [self.teacher,self.student,self.analyst,self.multi]:
            self.assertEqual(self.client_for(other).get('/api/users',{'q':'семенов','meta':'1'}).status_code,403)
        for params in [{'role':'owner'},{'state':'deleted'},{'graduation_year':'yesterday'},{'graduation_year':'2201'},{'meta':'1','page':'0'},{'q':'x'*161}]:
            self.assertEqual(client.get('/api/users',params).status_code,400)

    def test_people_admin_search_filters_before_limit_and_paginates(self):
        User.objects.bulk_create([User(email=f'page-{index}@plusfolio.example',name=f'Аккаунт {index:03}',name_search=f'аккаунт {index:03}',roles=['teacher']) for index in range(305)])
        client=self.client_for(self.admin)
        result=client.get('/api/users',{'meta':'1','q':'аккаунт','page':'11'})
        self.assertEqual(result.status_code,200);self.assertEqual(result.data['count'],305);self.assertEqual(len(result.data['results']),5)
        result=client.get('/api/users',{'q':'аккаунт 304'})
        self.assertEqual(len(result.data),1)
        self.assertEqual(result.data[0]['name'],'Аккаунт 304')

    def test_student_search_only_writers_active_students_and_minimal_fields(self):
        pupil=self._named_student();inactive=self._named_student('disabled@plusfolio.example');inactive.is_active=False;inactive.save()
        former=self._named_student('former@plusfolio.example');former.roles=['teacher'];former.save()
        path=f'/api/circles/{self.circle.pk}/student-search'
        teacher=self.client_for(self.teacher)
        self.assertEqual(teacher.get(path).data,[])
        result=teacher.get(path,{'q':'АНДРЕЕВИЧ семенов'})
        self.assertEqual(result.status_code,200);self.assertEqual([row['id'] for row in result.data],[pupil.student.pk])
        self.assertEqual(set(result.data[0]),{'id','public_id','user_public_id','name','graduation_year','current_grade','grade_label','email_hint','already_enrolled'})
        self.assertNotEqual(result.data[0]['email_hint'],pupil.email)
        self.assertFalse(result.data[0]['already_enrolled'])
        s.enroll(self.teacher,self.circle,student_id=pupil.student.pk)
        self.assertTrue(teacher.get(path,{'q':'семенов'}).data[0]['already_enrolled'])
        for account in [self.student,self.analyst,self.multi]:self.assertEqual(self.client_for(account).get(path,{'q':'семенов'}).status_code,403)
        self.assertEqual(teacher.get(f'/api/circles/{self.foreign.pk}/student-search',{'q':'семенов'}).status_code,404)
        self.assertEqual(APIClient().get(path,{'q':'семенов'}).status_code,403)
        self.assertEqual(teacher.get(path,{'q':'x'*161}).status_code,400)

    def test_student_search_limit_and_add_by_selected_id_keep_assignment_rules(self):
        for index in range(22):self._named_student(f'limit-{index}@plusfolio.example','Поисковый Иван Иванович')
        client=self.client_for(self.teacher);path=f'/api/circles/{self.circle.pk}'
        self.assertEqual(len(client.get(path+'/student-search',{'q':'поисковый'}).data),20)
        pupil=self._named_student()
        for _ in range(2):self.assertEqual(client.post(path+'/members',{'student_id':pupil.student.pk},format='json').status_code,200)
        self.assertEqual(pupil.student.assignments.count(),self.circle.issues.count())
        self.assertEqual(Enrollment.objects.filter(circle=self.circle,student=pupil.student,left_at=None).count(),1)
        self.assertEqual(client.post(path+'/members',{'student_id':True},format='json').status_code,400)
        self.assertEqual(client.post(path+'/members',{'student_id':pupil.student.pk,'email':pupil.email},format='json').status_code,400)
        self.assertEqual(self.client_for(self.analyst).post(path+'/members',{'student_id':pupil.student.pk},format='json').status_code,403)
        pupil.is_active=False;pupil.save()
        self.assertEqual(client.post(path+'/members',{'student_id':pupil.student.pk},format='json').status_code,400)

    def test_student_registration_requires_full_name_and_year_without_consuming_invite(self):
        invitation=s.invite(self.teacher,'identity@plusfolio.example','student',self.circle);token=self._invite_token(invitation)
        client=APIClient();url='/api/accept/'+token
        self.assertTrue(client.get(url).json()['student_profile_required'])
        self.assertEqual(client.post(url,[],format='json').status_code,400)
        base={'name':'Иванов Иван Иванович','graduation_year':2031,'password':'DemoPassword2026!'}
        for changes in [{'name':'Иванов И. И.'},{'name':'Иванов И И'},{'name':'Иванов Иван'},{'name':''},{'graduation_year':None},{'graduation_year':'2031'},{'graduation_year':True},{'graduation_year':2031.5},{'graduation_year':2201}]:
            self.assertEqual(client.post(url,{**base,**changes},format='json').status_code,400)
            self.assertFalse(User.objects.filter(email=invitation.email).exists())
            invitation.refresh_from_db();self.assertEqual(invitation.state,'pending')
        result=client.post(url,{**base,'name':' Иванов   Иван Иванович ','roles':['admin']},format='json')
        self.assertEqual(result.status_code,200);self.assertEqual(result.json()['user']['name'],'Иванов Иван Иванович')
        account=User.objects.get(email=invitation.email)
        self.assertEqual(account.roles,['student']);self.assertEqual(account.student.graduation_year,2031)
        self.assertTrue(result.json()['user']['profile_complete'])
        self.assertEqual(account.student.assignments.count(),self.circle.issues.count())

    def test_no_patronymic_is_explicit_and_other_roles_keep_normal_registration(self):
        invitation=s.invite(self.teacher,'no-patronymic@plusfolio.example','student',self.circle)
        account=s.accept_invitation(self._invite_token(invitation),{'name':'Ким Александр','graduation_year':2031,'no_patronymic':True,'password':'DemoPassword2026!'})
        self.assertTrue(account.student.no_patronymic)
        invitation=s.invite(self.admin,'new-teacher@plusfolio.example','teacher',self.circle)
        account=s.accept_invitation(self._invite_token(invitation),{'name':'Новый преподаватель','password':'DemoPassword2026!'})
        self.assertEqual(account.roles,['teacher']);self.assertFalse(hasattr(account,'student'))

    def test_existing_user_first_student_role_requires_identity_but_student_reinvite_does_not(self):
        invitation=s.invite(self.admin,self.teacher.email,'student',self.circle);token=self._invite_token(invitation)
        self.assertTrue(self.client_for(self.teacher).get('/api/accept/'+token).json()['student_profile_required'])
        with self.assertRaises(ValidationError):s.accept_invitation(token,{},self.teacher)
        self.teacher.refresh_from_db();self.assertNotIn('student',self.teacher.roles)
        account=s.accept_invitation(token,{'name':'Тестов Преподаватель Иванович','graduation_year':2031},self.teacher)
        self.assertEqual(set(account.roles),{'student','teacher'})
        self.assertEqual(account.student.graduation_year,2031)
        invitation=s.invite(self.admin,self.student.email,'student',self.foreign);token=self._invite_token(invitation)
        self.assertFalse(self.client_for(self.student).get('/api/accept/'+token).json()['student_profile_required'])
        account=s.accept_invitation(token,{},self.student)
        self.assertEqual(account.name,self.student.name)

    def test_student_account_cannot_replace_full_name_with_alias_or_clear_year_and_search_refreshes(self):
        pupil=self._named_student();client=self.client_for(pupil)
        self.assertEqual(client.patch('/api/account',['name','graduation_year','no_patronymic'],format='json').status_code,400)
        base={'name':pupil.name,'graduation_year':2031,'no_patronymic':False}
        for change in [{'name':'Вася'},{'graduation_year':None},{'no_patronymic':'false'},{'roles':['admin']},{'email':'new@plusfolio.example'}]:
            self.assertEqual(client.patch('/api/account',{**base,**change},format='json').status_code,400)
        result=client.patch('/api/account',{**base,'name':'Петров Алексей Андреевич','graduation_year':2032},format='json')
        self.assertEqual(result.status_code,200);self.assertTrue(result.data['user']['profile_complete']);self.assertEqual(result.data['user']['graduation_year'],2032)
        path=f'/api/circles/{self.circle.pk}/student-search'
        teacher=self.client_for(self.teacher)
        self.assertEqual(teacher.get(path,{'q':'семенов'}).data,[])
        self.assertEqual(teacher.get(path,{'q':'андреевич ПЕТРОВ'}).data[0]['id'],pupil.student.pk)
        self.assertEqual(AuditEvent.objects.filter(kind='account.student_profile',actor=pupil).count(),1)
        self.assertEqual(client.patch('/api/account',{**base,'name':'Петров Алексей Андреевич','graduation_year':2032},format='json').status_code,200)
        self.assertEqual(AuditEvent.objects.filter(kind='account.student_profile',actor=pupil).count(),1)



from django.test import SimpleTestCase

class SchoolGradeTests(SimpleTestCase):
    def test_school_year_rolls_over_on_september_first_and_not_in_january_or_june(self):
        from datetime import date
        from .school import education_info
        for day,grade in [(date(2026,8,31),6),(date(2026,9,1),7),(date(2026,12,31),7),(date(2027,1,1),7),(date(2027,6,30),7),(date(2027,8,31),7),(date(2027,9,1),8)]:
            with self.subTest(day=day):self.assertEqual(education_info(2031,day)['current_grade'],grade)

    def test_first_grade_last_grade_graduate_and_future_schoolchild(self):
        from datetime import date
        from .school import education_info
        for day,grade,label in [(date(2020,8,31),None,'Ещё не учится'),(date(2020,9,1),1,'1 класс'),(date(2031,8,31),11,'11 класс'),(date(2031,9,1),None,'Выпускник')]:
            with self.subTest(day=day):
                self.assertEqual(education_info(2031,day),{'graduation_year':2031,'current_grade':grade,'grade_label':label})
        self.assertEqual(education_info(None,date(2026,10,21)),{'graduation_year':None,'current_grade':None,'grade_label':'Класс не указан'})


from django.test import TransactionTestCase

@skipUnless(connection.vendor=='sqlite','Local demo backup format requires SQLite')
class BackupRuntimeTests(TransactionTestCase):
    def test_backup_restores_offline_tex_runtime_but_not_derived_pdf_cache(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from django.test import override_settings
        from .backup import create_bundle,restore_bundle
        with TemporaryDirectory() as directory:
            root=Path(directory)/'source';local=root/'.local';local.mkdir(parents=True)
            for name in ['tools/tectonic','tectonic-cache/bundle/file','tectonic-config/config.toml','pdf-cache/private.pdf']:
                path=local/name;path.parent.mkdir(parents=True,exist_ok=True);path.write_text('test runtime')
            (root/'source.py').write_text('# source')
            with override_settings(BASE_DIR=root,LOCAL_DIR=local):
                archive=Path(directory)/'test.pfb';manifest=create_bundle(archive)
                self.assertIn('.local/tools/tectonic',manifest['files'])
                self.assertIn('.local/tectonic-cache/bundle/file',manifest['files'])
                self.assertNotIn('.local/pdf-cache/private.pdf',manifest['files'])
                restored=Path(directory)/'restored';restore_bundle(archive,local/'backup.key',restored)
                self.assertTrue((restored/'.local/tools/tectonic').stat().st_mode&0o111)
                self.assertEqual((restored/'.local/tectonic-cache/bundle/file').read_text(),'test runtime')
