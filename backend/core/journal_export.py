"""Downloadable journal snapshots. User text is always written as literal text."""
import csv
import io
import re
from datetime import datetime
from django.core.exceptions import ValidationError
from django.http import HttpResponse
from django.utils import timezone
from django.utils.http import content_disposition_header
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN = '365B4A'
PALE = 'EAF1E3'
BLUE = 'EAF0F3'
ILLEGAL_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def columns(data):
    result = []
    for index, block in enumerate(data['blocks']):
        for number, problem in enumerate(block['problems'], 1):
            result.append((index, block, problem, f'№{number} · {problem["title"]}'))
        if data['show_group_stats']:
            result.append((index, block, None, 'Решено'))
    return result


def rows(data, cols):
    states = {(state['student_id'], state['problem_id']): state['solved'] for state in data['solves']}
    discussed = {state['problem_id']: state['discussed'] for state in data['discussions']}
    yield ['Разобрана в кружке'] + [('Да' if discussed.get(problem['id'], False) else 'Нет') if problem else '' for _, _, problem, _ in cols]
    for person in data['people']:
        row = [person['name']]
        for _, block, problem, _ in cols:
            if problem:
                row.append('Да' if states.get((person['id'], problem['id']), False) else 'Нет')
            else:
                row.append(sum(bool(states.get((person['id'], task['id']), False)) for task in block['problems']))
        yield row


def text(value):
    return ILLEGAL_XML.sub('', str(value))[:32767]


def tsv_cell(value):
    value = text(value)
    # Spreadsheet applications can execute formula-looking text on opening TSV.
    if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
        value = "'" + value
    return value


def xlsx(data, cols):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Журнал'
    sheet.sheet_view.showGridLines = False
    width = 1 + len(cols)
    # These immutable components are shared by cells in this workbook. Creating
    # equal Font/Border/Fill objects per cell makes style interning needlessly costly.
    body_font = Font(name='Calibri', size=11, color='273E34')
    block_font = Font(name='Calibri', size=11, bold=True, color=GREEN)
    task_font = Font(name='Calibri', size=10, bold=True, color=GREEN)
    left = Alignment(vertical='center', horizontal='left', wrap_text=True)
    centered = Alignment(vertical='center', horizontal='center', wrap_text=True)
    block_fills = (PatternFill('solid', fgColor=PALE), PatternFill('solid', fgColor=BLUE))
    discussion_fill = PatternFill('solid', fgColor='F0F3ED')
    solved_fill = PatternFill('solid', fgColor='E7F0DF')
    stripe_fill = PatternFill('solid', fgColor='FAFBF8')
    bottom_border = Border(bottom=Side(style='hair', color='DFE6DB'))
    def put(row, column, value):
        cell = sheet.cell(row, column)
        if isinstance(value, (int, float)):
            cell.value = value
        else:
            cell.value = text(value)
            cell.data_type = 's'
        cell.font = body_font
        cell.alignment = centered if column > 1 else left
        return cell
    def title_row(row, value):
        put(row, 1, value)
        if width > 1:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    title_row(1, f'Plusfolio · {data["circle"]["name"]}')
    sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color=GREEN)
    sheet.row_dimensions[1].height = 32
    title_row(2, f'Журнал сдачи · Снимок на {timezone.localtime():%d.%m.%Y %H:%M} · Листков: {len(data["blocks"])} · Школьников: {len(data["people"])}')
    title_row(3, 'Да / Нет — задача решена. Решённость общая между кружками; разбор относится к этому кружку.')
    sheet.row_dimensions[2].height = 25
    sheet.row_dimensions[3].height = 30
    put(5, 1, 'Школьник')
    sheet.merge_cells('A5:A6')
    start = 2
    for index, block in enumerate(data['blocks']):
        count = sum(1 for block_index, *_ in cols if block_index == index)
        if not count:
            continue
        issued_on = datetime.fromisoformat(block['issued_on']).strftime('%d.%m.%Y')
        put(5, start, f'{block["worksheet"]["title"]} · {issued_on}')
        if count > 1:
            sheet.merge_cells(start_row=5, start_column=start, end_row=5, end_column=start + count - 1)
        for column in range(start, start + count):
            cell = sheet.cell(5, column)
            cell.fill = block_fills[index % 2]
            cell.font = block_font
            cell.alignment = centered
        start += count
    for column, (index, _, _, label) in enumerate(cols, 2):
        cell = put(6, column, label)
        cell.fill = block_fills[index % 2]
        cell.font = task_font
    for row_number, values in enumerate(rows(data, cols), 7):
        sheet.row_dimensions[row_number].height = 28
        for column, value in enumerate(values, 1):
            cell = put(row_number, column, value)
            if row_number == 7:
                cell.fill = discussion_fill
            elif column > 1 and value == 'Да':
                cell.fill = solved_fill
            elif row_number % 2 == 0:
                cell.fill = stripe_fill
            cell.border = bottom_border
    sheet['A5'].fill = PatternFill('solid', fgColor=GREEN)
    sheet['A5'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    sheet.row_dimensions[5].height = 42
    sheet.row_dimensions[6].height = max(64, min(160, max((len(label) // 18 + 1) * 12 for *_, label in cols))) if cols else 40
    sheet.column_dimensions['A'].width = 30
    for column, (_, _, problem, _) in enumerate(cols, 2):
        sheet.column_dimensions[get_column_letter(column)].width = 19 if problem else 12
    sheet.freeze_panes = 'B7'
    sheet.print_title_rows = '1:7'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 0
    sheet.page_setup.fitToHeight = 0
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export(data, file_format):
    if file_format not in {'xlsx', 'tsv'}:
        raise ValidationError('Доступны форматы Excel (.xlsx) и TSV (.tsv).')
    cols = columns(data)
    if len(cols) + 1 > 16384 or (len(data['people']) + 8) * (len(cols) + 1) > 1_000_000:
        raise ValidationError('Выгрузка слишком большая. Выберите меньше листков или уточните поиск школьника.')
    if file_format == 'xlsx':
        body = xlsx(data, cols)
        content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        buffer = io.BytesIO()
        with io.TextIOWrapper(buffer, encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream, delimiter='\t', lineterminator='\r\n')
            writer.writerow(['Школьник'] + [tsv_cell(f'{block["worksheet"]["title"]} · {block["issued_on"]}') for _, block, _, _ in cols])
            writer.writerow([''] + [tsv_cell(label) for _, _, _, label in cols])
            writer.writerows([tsv_cell(value) for value in row] for row in rows(data, cols))
            stream.flush()
            body = buffer.getvalue()
        content_type = 'text/tab-separated-values; charset=utf-8'
    response = HttpResponse(body, content_type=content_type)
    filename = f'plusfolio-{data["circle"]["id"][:8]}-{timezone.localdate():%Y%m%d}.{file_format}'
    response['Content-Disposition'] = content_disposition_header(True, filename)
    response['Cache-Control'] = 'no-store'
    return response
